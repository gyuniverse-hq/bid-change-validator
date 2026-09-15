"""Local evaluation adapter: read workbook cells without editing/calculating Excel.

Does not add XLSX support to the production collector. Run with the bundled
document Python runtime. Original snapshot and workbook remain unchanged.
"""
import argparse
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import openpyxl


def render(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else str(value)


def prepare(snapshot, originals, output):
    if output.exists() or output.resolve() == snapshot.resolve():
        raise ValueError('Output must be new')
    data = json.loads(snapshot.read_text(encoding='utf-8'))
    changes = []
    for document in data['documents']:
        if not document['name'].lower().endswith('.xlsx'):
            continue
        digest = document['file_sha256']
        if len(digest) != 64 or any(c not in '0123456789abcdef' for c in digest):
            raise ValueError('Invalid source digest')
        path = originals / (digest + '.xlsx')
        if hashlib.sha256(path.read_bytes()).hexdigest() != digest:
            raise ValueError('Workbook hash mismatch')
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
        cached = openpyxl.load_workbook(path, read_only=True, data_only=True)
        blocks, formulas = [], 0
        sheets = []
        try:
            for sheet in workbook:
                sheets.append({'name': sheet.title, 'rows': sheet.max_row, 'columns': sheet.max_column})
                cached_rows = cached[sheet.title].iter_rows()
                for row_number, (row, values) in enumerate(zip(sheet.iter_rows(), cached_rows, strict=True), start=1):
                    cells = []
                    for cell, value_cell in zip(row, values, strict=True):
                        if cell.value is None or cell.value == '':
                            continue
                        if cell.data_type == 'f':
                            formulas += 1
                            value = '수식 ' + str(cell.value) + '; 파일에 저장된 계산값(재계산 미검증): ' + render(value_cell.value)
                        else:
                            value = render(cell.value)
                        cells.append(cell.coordinate + ': ' + value)
                    if cells:
                        blocks.append({'block_index': len(blocks), 'location': f'{sheet.title}!{row_number}',
                            'text': sheet.title + ' — ' + ' / '.join(cells)})
        finally:
            workbook.close()
            cached.close()
        text = '\n\n'.join(b['text'] for b in blocks)
        new_hash = hashlib.sha256(text.encode()).hexdigest()
        document.update(extracted_text=text, extracted_blocks=blocks, extracted_char_count=len(text),
            extracted_text_sha256=new_hash, extraction_status='EXTRACTED', extraction_error=None,
            text_extractor='LOCAL_XLSX_CELL_REVIEW_OPENPYXL')
        changes.append({'document_id': document['id'], 'source_sha256': digest,
                        'extracted_sha256': new_hash, 'sheets': sheets, 'formula_cells': formulas,
                        'scope': 'cell text/formulas/stored values only; no workbook editing, recalculation, or image interpretation'})
    if not changes:
        raise ValueError('No XLSX documents')
    output.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    output.with_suffix('.xlsx-provenance.json').write_text(json.dumps({'changes': changes,
        'input_snapshot_sha256': hashlib.sha256(snapshot.read_bytes()).hexdigest(),
        'output_sha256': hashlib.sha256(output.read_bytes()).hexdigest()}, ensure_ascii=False, indent=2), encoding='utf-8')
    print('Prepared local cell extraction:', len(changes), 'documents')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--snapshot', type=Path, required=True)
    parser.add_argument('--originals', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    prepare(args.snapshot, args.originals, args.output)
